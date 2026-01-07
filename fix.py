import openpyxl
import pyrebase
import time
config = {
  "apiKey": "AIzaSyBGkQh9ygRa6-iEY2V4mWQ8YGemFZA4-HI",
  "authDomain": "thkstorage-ec701.firebaseapp.com",
  "databaseURL": "https://thkstorage-ec701-default-rtdb.asia-southeast1.firebasedatabase.app",
  "storageBucket": "thkstorage-ec701.firebasestorage.app"
}

firebase = pyrebase.initialize_app(config)
db = firebase.database()
wb = openpyxl.load_workbook('name_list.xlsx') 
for k in wb.sheetnames:
  sheet = wb[k] 
  # เข้าถึงค่าเซลล์
  n=0
  nm = ""
  print(f"row {(sheet.max_row)}")
  for i in sheet.rows:
      n+=1
      cell_value = sheet.cell(n,3).value
      class_m = sheet.cell(n,2).value
      no_studio = sheet.cell(n,1).value
      if class_m is not None:
        if type(class_m) != type(1):
          if "ชั้นมัธยมศึกษาปีที่ " in class_m:
            nm = class_m.replace("ชั้นมัธยมศึกษาปีที่","").replace(" ","").split("ปีการศึกษา")[0]
            print('class',nm)
      if cell_value is None or "ชื่อ - นามสกุล" in cell_value or "เลขประจำตัว" in cell_value or len(cell_value.strip()) < 0:
          continue
      if type(no_studio) != type(1) and no_studio is not None:
        if "เลขที่" in no_studio:
          continue
      first_name = sheet.cell(n,3).value
      last_name = ""
      if sheet.cell(n,4).value is not None:
        last_name = sheet.cell(n,4).value
      try:
        last_name = sheet.cell(n,4).value + (" "+sheet.cell(n,5).value if sheet.cell(n,5).value is not None else "")
      except:
        pass
      number_studio = sheet.cell(n,2).value
      print(f"{nm}/{no_studio}",n,number_studio, first_name, last_name)
      data = {}
      data["data"] = first_name+" "+last_name
      db.child(f"student_2026/{nm}/{no_studio}/{number_studio}").set(data)
      
      # time.sleep(1)
      
